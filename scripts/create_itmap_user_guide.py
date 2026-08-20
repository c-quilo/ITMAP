from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ITMAP_User_Guide.docx"
SWOOSH = ROOT / "public" / "scs-swoosh.png"
LOGO = ROOT / "src" / "assets" / "imperial-logo.png"

INK = RGBColor(15, 31, 45)
MUTED = RGBColor(86, 98, 110)
BLUE = RGBColor(32, 88, 124)
TEAL = RGBColor(20, 128, 122)
GOLD = RGBColor(166, 119, 25)
LIGHT_BLUE = "EAF3F8"
LIGHT_TEAL = "E9F6F4"
LIGHT_GOLD = "FFF6DF"
LIGHT_GRAY = "F3F5F7"
BORDER = "D7DEE6"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=130, bottom=90, end=130):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def set_run(run, size=None, color=None, bold=None, italic=None, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def paragraph_border_bottom(paragraph, color="2E74B5", size="12", space="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 7),
        ("Heading 2", 13, BLUE, 11, 5),
        ("Heading 3", 11.5, RGBColor(31, 77, 120), 8, 3),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_header_footer(doc):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("ITMAP User Guide")
    set_run(r, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Imperial Thematic Mapping and Profiling tool")
    set_run(r, size=8.5, color=MUTED)


def add_cover(doc):
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run().add_picture(str(LOGO), width=Inches(2.05))

    doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kr = kicker.add_run("FIELD GUIDE")
    set_run(kr, size=10, color=GOLD, bold=True)
    kicker.paragraph_format.space_after = Pt(8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("ITMAP")
    set_run(tr, size=34, color=INK, bold=True, font="Calibri")
    title.paragraph_format.space_after = Pt(2)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("How to find, compare, save, and explain relevant Imperial researchers")
    set_run(sr, size=14, color=MUTED)
    sub.paragraph_format.space_after = Pt(16)

    if SWOOSH.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(SWOOSH), width=Inches(4.9))
        p.paragraph_format.space_after = Pt(10)

    meta = doc.add_table(rows=1, cols=3)
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_width(meta, [2.05, 2.05, 2.05])
    values = [
        ("Best for", "Mission-led expert discovery"),
        ("Audience", "Internal users and collaborators"),
        ("Last updated", "July 2026"),
    ]
    for idx, (label, value) in enumerate(values):
        cell = meta.cell(0, idx)
        set_cell_shading(cell, LIGHT_BLUE if idx != 1 else LIGHT_TEAL)
        set_cell_border(cell, "FFFFFF")
        set_cell_margins(cell, 120, 160, 120, 160)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label.upper())
        set_run(r, size=7.5, color=MUTED, bold=True)
        p.add_run("\n")
        r = p.add_run(value)
        set_run(r, size=10.5, color=INK, bold=True)

    doc.add_paragraph()
    lead = doc.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lr = lead.add_run(
        "ITMAP turns a mission statement into a ranked, explainable shortlist of Imperial researchers, "
        "using profiles, papers, external evidence, School Missions, saved lists, and CSV export."
    )
    set_run(lr, size=11.5, color=INK)
    lead.paragraph_format.space_after = Pt(8)
    paragraph_border_bottom(lead, color="14807A", size="10", space="10")

    doc.add_page_break()


def add_callout(doc, title, body, fill=LIGHT_TEAL):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, "FFFFFF")
    set_cell_margins(cell, 130, 170, 130, 170)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    set_run(r, size=10, color=BLUE, bold=True)
    p.add_run("\n")
    r = p.add_run(body)
    set_run(r, size=10.5, color=INK)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.14)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run(r, size=10.5, color=INK)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.14)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run(r, size=10.5, color=INK)


def add_feature_grid(doc):
    rows = [
        ("ITMAP rerank", "Re-checks the narrowed pool using profile, role, paper titles, and evidence. Keep it on for most mission searches."),
        ("Media, grants and startups", "Adds public-web, video, startup/spinout, and UKRI grant signals. Useful for impact evidence."),
        ("Generate Summary", "Creates a top-of-results synthesis: themes, notable researchers, and gaps. Cached for repeat use."),
        ("Check Missions", "Matches top results against School Missions and adds labels such as Sustainability · Re-Engineering."),
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_width(table, [1.9, 4.6])
    for cell, text in zip(table.rows[0].cells, ("Control", "What it does")):
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        r = cell.paragraphs[0].add_run(text)
        set_run(r, size=9.5, color=BLUE, bold=True)
    for name, detail in rows:
        cells = table.add_row().cells
        for c in cells:
            set_cell_border(c)
            set_cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r = cells[0].paragraphs[0].add_run(name)
        set_run(r, size=10, color=INK, bold=True)
        r = cells[1].paragraphs[0].add_run(detail)
        set_run(r, size=9.6, color=INK)


def add_result_anatomy(doc):
    table = doc.add_table(rows=1, cols=3)
    set_table_width(table, [1.65, 2.3, 2.55])
    headers = ("Part of result", "What to look for", "How to use it")
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT_TEAL)
        set_cell_border(cell)
        set_cell_margins(cell)
        r = cell.paragraphs[0].add_run(text)
        set_run(r, size=9.2, color=BLUE, bold=True)
    rows = [
        ("Match label", "Strong, Moderate, or Weak", "Use it as a quick triage signal, then read the explanation."),
        ("Why they matched", "Always visible on each card", "Explains how the profile, role, and publications connect to the query."),
        ("Relevant publications", "Up to 10 paper titles", "Open the list to inspect whether the evidence really matches the mission."),
        ("External signals", "Media, video, startup, grant", "Look for impact evidence when translational relevance matters."),
        ("Mission label", "School theme and mission", "Appears after Check Missions and can be used as a filter."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            set_cell_border(cells[idx])
            set_cell_margins(cells[idx])
            r = cells[idx].paragraphs[0].add_run(text)
            set_run(r, size=9.4 if idx else 9.6, color=INK, bold=(idx == 0))


def add_workflow_strip(doc):
    table = doc.add_table(rows=1, cols=5)
    set_table_width(table, [1.3, 1.3, 1.3, 1.3, 1.3])
    steps = [
        ("1", "Describe mission"),
        ("2", "Search profiles + papers"),
        ("3", "Review explanations"),
        ("4", "Filter + save"),
        ("5", "Export or share"),
    ]
    fills = [LIGHT_BLUE, LIGHT_TEAL, LIGHT_GOLD, LIGHT_BLUE, LIGHT_TEAL]
    for idx, (num, text) in enumerate(steps):
        cell = table.cell(0, idx)
        set_cell_shading(cell, fills[idx])
        set_cell_border(cell, "FFFFFF")
        set_cell_margins(cell, 130, 100, 130, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(num)
        set_run(r, size=14, color=BLUE, bold=True)
        p.add_run("\n")
        r = p.add_run(text)
        set_run(r, size=9.2, color=INK, bold=True)


def build():
    doc = Document()
    style_document(doc)
    add_header_footer(doc)
    add_cover(doc)

    doc.add_heading("The 60-second version", level=1)
    add_callout(
        doc,
        "Use ITMAP like a research scout.",
        "Write a mission, let ITMAP find a ranked pool, read why each person matched, then save the people you want to contact. "
        "Generate a pool summary when you need an overview, and Check Missions when you want to align people with School of Convergence Science missions.",
        LIGHT_TEAL,
    )
    doc.add_paragraph()
    add_workflow_strip(doc)

    doc.add_heading("Start a search", level=1)
    doc.add_heading("Semantic search", level=2)
    p = doc.add_paragraph()
    p.add_run("Use this for missions, questions, and fuzzy intent. ").bold = True
    p.add_run("Examples: “AI for weather”, “sustainable textiles”, “environmental exposure”, or a pasted mission brief.")
    add_numbered(doc, [
        "Open the Search tab.",
        "Keep Semantic selected.",
        "Type the mission in natural language, or attach a text file.",
        "Choose whether to keep ITMAP rerank and Media, grants and startups enabled.",
        "Click Search Researchers.",
    ])

    doc.add_heading("Keyword search", level=2)
    p = doc.add_paragraph(
        "Use keyword search when you need exact terms to appear in profiles, fields, or papers. It is stricter than semantic search and is best for targeted follow-up."
    )
    set_run(p.runs[0], size=10.5, color=INK)

    doc.add_page_break()
    doc.add_heading("Search controls", level=1)
    add_feature_grid(doc)

    doc.add_heading("Read the results", level=1)
    add_result_anatomy(doc)
    add_callout(
        doc,
        "Important habit",
        "Do not judge a match by the label alone. Read “Why they matched” and open Relevant Publications. A strong domain match is not always a strong method+domain match.",
        LIGHT_GOLD,
    )

    doc.add_page_break()
    doc.add_heading("Use the top-of-results actions", level=1)
    doc.add_heading("Generate Summary", level=2)
    add_bullets(doc, [
        "Creates a short synthesis of the current researcher pool.",
        "Highlights themes, notable researchers, and caveats or gaps.",
        "Appears above the result cards and is cached for the same query and top-result set.",
        "Use it when sharing findings with someone who does not want to inspect every card.",
    ])
    doc.add_heading("Check Missions", level=2)
    add_bullets(doc, [
        "Runs an extra pass against the School of Convergence Science mission brief.",
        "Adds labels such as Human and Artificial Intelligence · SYMBIOSIS or Sustainability · Re-Engineering.",
        "Unlocks School Missions filters in the sidebar.",
        "Use it after you have a good candidate pool, not before the first search.",
    ])

    doc.add_heading("Filter the pool", level=1)
    add_bullets(doc, [
        "Match Strength filters by Strong, Moderate, or Weak.",
        "Grade / Role filters by seniority terms such as Professor, Reader, Lecturer, or Research Fellow.",
        "Keywords are generated from the current results, including profiles, fields, and publication titles.",
        "Faculty and Department filters are generated from the current search results.",
        "School Missions filters appear only after you run Check Missions.",
    ])

    doc.add_heading("Save and export researchers", level=1)
    add_bullets(doc, [
        "Click the bookmark icon on any researcher card to save that person.",
        "Open the Saved tab to review your saved list.",
        "Each saved researcher keeps the search that produced it, so you can trace why they were saved.",
        "Click Export CSV to download the saved list for sharing or follow-up.",
    ])

    doc.add_heading("Use the Graph tab", level=1)
    add_bullets(doc, [
        "Use Graph after a search to inspect co-authorship and adjacent connections across the returned pool.",
        "Bridge matches are useful when a person connects otherwise separate clusters.",
        "Adjacent matches can be useful collaborators, but should be checked carefully against the mission.",
    ])

    doc.add_heading("Good search patterns", level=1)
    table = doc.add_table(rows=1, cols=2)
    set_table_width(table, [2.2, 4.3])
    for cell, text in zip(table.rows[0].cells, ("Pattern", "Example")):
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        r = cell.paragraphs[0].add_run(text)
        set_run(r, size=9.5, color=BLUE, bold=True)
    rows = [
        ("Method + domain", "machine learning for wildfires; AI for weather forecasting"),
        ("Problem + population", "air pollution exposure in children; resilient water systems in mountain regions"),
        ("Material + application", "sustainable textiles for circular fashion; biomass-derived carbon materials"),
        ("Policy + system", "energy access for low-income communities; circular resource stewardship"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            set_cell_border(cells[idx])
            set_cell_margins(cells[idx])
            r = cells[idx].paragraphs[0].add_run(text)
            set_run(r, size=9.8, color=INK, bold=(idx == 0))

    doc.add_heading("Troubleshooting", level=1)
    add_bullets(doc, [
        "If results feel too broad, add both the method and the application area.",
        "If a relevant person is missing, try a synonym or a paper-title phrase.",
        "If papers look irrelevant, open Relevant Publications and check whether the person matched mostly by profile.",
        "If external evidence looks noisy, rerun with Media, grants and startups turned off.",
        "If School Missions labels are absent, run Check Missions after results appear.",
    ])

    doc.add_heading("What ITMAP does not replace", level=1)
    add_callout(
        doc,
        "Human judgement still matters.",
        "ITMAP is a discovery and triage tool. Before contacting people or making decisions, inspect the profile, publications, and explanations, then confirm fit with the researcher or department.",
        LIGHT_BLUE,
    )

    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
